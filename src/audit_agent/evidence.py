"""Evidence loading helpers shared by the controls.

These utilities keep cell-level provenance (sheet + cell reference) so that
every reconciliation finding can cite the exact source it came from, and they
resolve sheets/columns by fuzzy header match against a required-column
manifest rather than hard-coding sheet names or column order.
"""

from __future__ import annotations

import hashlib
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


class MissingColumnError(ValueError):
    """Raised when a required logical column cannot be resolved in a sheet."""


def normalise_email(value: object) -> str:
    """Lower-case, strip, and NFKC-normalise an email for use as a join key."""
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip().lower()


def canon(text: object) -> str:
    """Canonical form of a header/label: lowercase, alphanumeric only."""
    return re.sub(r"[^a-z0-9]", "", str(text).strip().lower()) if text is not None else ""


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


# Excel's day-zero is 1899-12-30 (the 1900 leap-year bug is only material for
# dates before 1900-03-01, which do not occur in this domain).
_EXCEL_EPOCH = datetime(1899, 12, 30)


def coerce_date(value: object) -> date | None:
    """Coerce an Excel cell value to a ``date``.

    openpyxl returns ``datetime`` for date-formatted cells; some exports store
    the raw serial number or an ISO string instead, so handle all three.
    """
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return (_EXCEL_EPOCH + _timedelta_days(float(value))).date()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _timedelta_days(days: float):
    from datetime import timedelta

    return timedelta(days=days)


def find_sheet_name(sheet_names: list[str], candidates: list[str]) -> str | None:
    """Resolve a sheet by fuzzy match against candidate names."""
    wanted = [canon(c) for c in candidates]
    for name in sheet_names:
        cname = canon(name)
        if any(cname == w or w in cname or cname in w for w in wanted):
            return name
    return None


def resolve_columns(header: list[object], manifest: dict[str, list[str]]) -> dict[str, int]:
    """Map each logical column name to a header index via fuzzy matching.

    ``manifest`` maps a logical name to a list of acceptable header aliases.
    Raises :class:`MissingColumnError` listing any logical names not found.
    """
    canon_header = [canon(h) for h in header]
    resolved: dict[str, int] = {}
    missing: list[str] = []
    for logical, aliases in manifest.items():
        idx = _match_column(canon_header, aliases)
        if idx is None:
            missing.append(logical)
        else:
            resolved[logical] = idx
    if missing:
        raise MissingColumnError(
            f"missing required columns {missing}; header was {list(header)}"
        )
    return resolved


def _match_column(canon_header: list[str], aliases: list[str]) -> int | None:
    canon_aliases = [canon(a) for a in aliases]
    # Exact canonical match first, then substring containment.
    for idx, h in enumerate(canon_header):
        if h and h in canon_aliases:
            return idx
    for idx, h in enumerate(canon_header):
        if h and any(a and (a in h or h in a) for a in canon_aliases):
            return idx
    return None


class Table:
    """A loaded sheet with logical columns resolved and rows kept positional.

    Each record is a dict of ``logical_name -> value`` plus ``_row`` (the
    1-based Excel row number) so findings can cite an exact cell.
    """

    def __init__(self, file_name: str, sheet: str, col_letters: dict[str, str],
                 records: list[dict]):
        self.file_name = file_name
        self.sheet = sheet
        self.col_letters = col_letters
        self.records = records

    def cell_ref(self, logical: str, record: dict) -> str:
        letter = self.col_letters.get(logical, "?")
        return f"{self.file_name}!{self.sheet}!{letter}{record['_row']}"


def load_table(path: str | Path, sheet_candidates: list[str],
               manifest: dict[str, list[str]], header_scan_rows: int = 12) -> Table:
    """Load a sheet into a :class:`Table`, resolving sheet and columns fuzzily."""
    path = Path(path)
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet_name = find_sheet_name(wb.sheetnames, sheet_candidates)
    if sheet_name is None:
        raise MissingColumnError(
            f"no sheet matching {sheet_candidates} in {path.name} "
            f"(have {wb.sheetnames})"
        )
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))

    header_idx, col_map = _detect_header(rows[:header_scan_rows], manifest)
    col_letters = {logical: get_column_letter(idx + 1) for logical, idx in col_map.items()}

    records: list[dict] = []
    for offset, raw in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(c not in (None, "") for c in raw):
            continue
        record = {logical: _cell(raw, idx) for logical, idx in col_map.items()}
        record["_row"] = offset
        records.append(record)
    wb.close()
    return Table(path.name, sheet_name, col_letters, records)


def _cell(raw: tuple, idx: int):
    return raw[idx] if idx < len(raw) else None


def _detect_header(rows: list[tuple], manifest: dict[str, list[str]]) -> tuple[int, dict[str, int]]:
    """Find the header row that resolves the most manifest columns."""
    best_idx, best_map, best_score = None, None, -1
    for idx, raw in enumerate(rows):
        try:
            col_map = resolve_columns(list(raw), manifest)
        except MissingColumnError:
            continue
        score = len(col_map)
        if score > best_score:
            best_idx, best_map, best_score = idx, col_map, score
    if best_map is None:
        # Re-run on the most-populated row to surface a precise error.
        densest = max(range(len(rows)), key=lambda i: sum(1 for c in rows[i] if c not in (None, "")))
        resolve_columns(list(rows[densest]), manifest)  # raises MissingColumnError
    return best_idx, best_map


def load_key_value(path: str | Path, sheet_candidates: list[str]) -> dict[str, tuple]:
    """Load a two-column key/value metadata sheet (e.g. a Cover sheet).

    Returns ``{key: (value, cell_ref)}`` so values can be cited.
    """
    path = Path(path)
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet_name = find_sheet_name(wb.sheetnames, sheet_candidates)
    if sheet_name is None:
        wb.close()
        return {}
    ws = wb[sheet_name]
    out: dict[str, tuple] = {}
    for row_idx, raw in enumerate(ws.iter_rows(values_only=True), start=1):
        if not raw or raw[0] in (None, ""):
            continue
        key = str(raw[0]).strip().rstrip(":")
        value = raw[1] if len(raw) > 1 else None
        out[key] = (value, f"{path.name}!{sheet_name}!B{row_idx}")
    wb.close()
    return out
