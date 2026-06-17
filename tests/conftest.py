"""Shared fixtures: real-data paths and a synthetic-workbook builder."""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
SRC = REPO_ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

DATA_DIR = REPO_ROOT / "data"
UAR_REAL = DATA_DIR / "user-access-review"
ICR_REAL = DATA_DIR / "independent-code-review"

ACCESS_HEADERS = ["Username", "Full Name", "Email", "Role", "Account Status",
                  "MFA Enabled", "Last Login", "Date Provisioned"]
REVIEW_HEADERS = ["Email", "Name", "Role", "Account Status", "Reviewer Decision",
                  "Reviewer Comment", "Reviewed By"]
HRIS_HEADERS = ["Employee ID", "First Name", "Last Name", "Work Email", "Department",
                "Job Title", "Employment Status", "Hire Date", "Termination Date"]

DEFAULT_COVER = {
    "Review Period": "Q2 2026",
    "Review Type": "Periodic (Quarterly) User Access Review",
    "Review Completed": date(2026, 6, 30),
    "Reviewer / Approver": "Priya Nadkarni, Director, Finance Systems (System Owner)",
    "Reviewer Sign-off:": "Approved electronically by Priya Nadkarni on 2026-06-30.",
}


@pytest.fixture
def make_sample(tmp_path):
    """Build a synthetic UAR sample folder and return its path.

    access/review/hris are lists of dicts keyed by the human column labels.
    """

    def _build(access, review, hris, cover=None):
        cover = cover or DEFAULT_COVER
        uar = Workbook()
        cov_ws = uar.active
        cov_ws.title = "Cover"
        cov_ws.append(["User Access Review", ""])
        for k, v in cover.items():
            cov_ws.append([k, v])
        _sheet(uar.create_sheet("System Access Export"), ACCESS_HEADERS, access)
        _sheet(uar.create_sheet("Access Review"), REVIEW_HEADERS, review)
        # plausible extra sheet to exercise fuzzy sheet resolution
        uar.create_sheet("Summary & Observations").append(["Summary", ""])
        uar_path = tmp_path / "uar-netsuite-q2-2026.xlsx"
        uar.save(uar_path)

        hwb = Workbook()
        _sheet(hwb.active, HRIS_HEADERS, hris)
        hwb.active.title = "Employees"
        hris_path = tmp_path / "hris-employee-export.xlsx"
        hwb.save(hris_path)
        return tmp_path

    return _build


def _sheet(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
